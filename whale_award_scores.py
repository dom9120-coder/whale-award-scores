#!/usr/bin/env python3
"""구글폼 응답 엑셀에서 개인정보 고래상 평가 점수와 순위를 계산한다."""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CANDIDATES = tuple("abcdefghi")

EVALUATOR_WEIGHTS = {
    "차관님": 0.3,
    "사무처장님": 0.2,
    "대변인님": 0.1,
    "기획조정관님": 0.1,
    "정책국장님": 0.1,
    "조사국장님": 0.1,
    "예방조정심의관님": 0.1,
    "마이데이터 추진단장님(대리)": 0.1,
}

EXCLUSIONS = {
    "차관님": set(),
    "사무처장님": set(),
    "대변인님": set(),
    "기획조정관님": {"a", "b"},
    "정책국장님": {"c", "d"},
    "조사국장님": {"e", "f"},
    "예방조정심의관님": {"g", "h"},
    "마이데이터 추진단장님(대리)": {"i"},
}

# 구글폼이나 안내문에서 사용될 수 있는 직책 표기 차이를 흡수한다.
ROLE_ALIASES = {
    "기조관님": "기획조정관님",
    "사전예방심의관님": "예방조정심의관님",
    "예방조정심의관님": "예방조정심의관님",
    "마이데이터추진단장님(대리)": "마이데이터 추진단장님(대리)",
}

HEADER_PATTERN = re.compile(
    r"^([a-i])\s+\[(.+?)\((\d+)\s*/\s*(\d+)\s*/\s*(\d+)\)\](?:\s+\d+)?$"
)
LEVEL_INDEX = {"상": 0, "중": 1, "하": 2}

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
SUBTITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
WHITE_BOLD = Font(color="FFFFFF", bold=True)


def normalize_role(value: object) -> str:
    role = str(value or "").strip()
    return ROLE_ALIASES.get(role, role)


def parse_timestamp(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if value is None:
        return datetime.min
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return datetime.min


def read_latest_responses(input_path: Path) -> tuple[list[object], dict[str, tuple[object, ...]]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("응답 시트가 비어 있습니다.")

    headers = list(rows[0])
    latest: dict[str, tuple[object, ...]] = {}
    latest_time: dict[str, datetime] = {}

    for row in rows[1:]:
        role = normalize_role(row[1] if len(row) > 1 else None)
        if not role:
            continue
        timestamp = parse_timestamp(row[0] if row else None)
        if role not in latest or timestamp >= latest_time[role]:
            latest[role] = row
            latest_time[role] = timestamp

    return headers, latest


def calculate_scores(
    headers: list[object], responses: dict[str, tuple[object, ...]]
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[str]]:
    errors: list[str] = []
    details: list[dict[str, object]] = []

    missing_roles = set(EVALUATOR_WEIGHTS) - set(responses)
    unknown_roles = set(responses) - set(EVALUATOR_WEIGHTS)
    if missing_roles:
        errors.append("응답이 없는 평가자: " + ", ".join(sorted(missing_roles)))
    if unknown_roles:
        errors.append("알 수 없는 평가자: " + ", ".join(sorted(unknown_roles)))

    parsed_headers: dict[int, tuple[str, str, tuple[int, int, int]]] = {}
    for index, header in enumerate(headers[2:], start=2):
        match = HEADER_PATTERN.match(str(header or "").strip())
        if not match:
            errors.append(f"해석할 수 없는 평가 항목 헤더: {header!r}")
            continue
        candidate, item, high, middle, low = match.groups()
        parsed_headers[index] = (candidate, item.strip(), (int(high), int(middle), int(low)))

    evaluator_candidate_scores: dict[tuple[str, str], int] = defaultdict(int)
    evaluator_candidate_counts: dict[tuple[str, str], int] = defaultdict(int)

    for role, row in responses.items():
        if role not in EVALUATOR_WEIGHTS:
            continue
        seen_items: set[tuple[str, str]] = set()
        for column_index, (candidate, item, points) in parsed_headers.items():
            value = row[column_index] if column_index < len(row) else None
            if value is None or str(value).strip() == "":
                continue
            level = str(value).strip()
            if level not in LEVEL_INDEX:
                errors.append(f"{role} / {candidate} / {item}: 알 수 없는 평가값 {value!r}")
                continue
            if candidate in EXCLUSIONS[role]:
                errors.append(f"{role}가 상피 대상 {candidate}를 평가했습니다.")
                continue
            item_key = (candidate, item)
            if item_key in seen_items:
                errors.append(f"{role} / {candidate} / {item}: 중복 응답이 있습니다.")
                continue
            seen_items.add(item_key)
            score = points[LEVEL_INDEX[level]]
            evaluator_candidate_scores[(role, candidate)] += score
            evaluator_candidate_counts[(role, candidate)] += 1

        for candidate in CANDIDATES:
            count = evaluator_candidate_counts[(role, candidate)]
            if candidate in EXCLUSIONS[role]:
                if count:
                    errors.append(f"{role} / {candidate}: 상피 대상인데 {count}개 항목이 입력됐습니다.")
            elif count != 5:
                errors.append(f"{role} / {candidate}: 평가 항목이 {count}개입니다(정상값 5개).")

    for candidate in CANDIDATES:
        for role, weight in EVALUATOR_WEIGHTS.items():
            excluded = candidate in EXCLUSIONS[role]
            raw_score = evaluator_candidate_scores[(role, candidate)] if not excluded else None
            details.append(
                {
                    "피평가자": candidate,
                    "평가자": role,
                    "상피여부": "제외" if excluded else "",
                    "원점수": raw_score,
                    "가중치": weight,
                    "가중점수": None if raw_score is None else raw_score * weight,
                    "항목수": evaluator_candidate_counts[(role, candidate)],
                }
            )

    summaries: list[dict[str, object]] = []
    for candidate in CANDIDATES:
        candidate_details = [d for d in details if d["피평가자"] == candidate and not d["상피여부"]]
        weight_sum = sum(float(d["가중치"]) for d in candidate_details)
        weighted_score = sum(float(d["가중점수"] or 0) for d in candidate_details)
        if abs(weight_sum - 1.0) > 1e-9:
            errors.append(f"{candidate}: 유효 평가자 가중치 합이 {weight_sum:.3f}입니다(정상값 1.000).")
        summaries.append(
            {
                "피평가자": candidate,
                "최종점수": weighted_score,
                "유효 평가자 수": len(candidate_details),
                "가중치 합": weight_sum,
            }
        )

    ordered_scores = sorted((float(s["최종점수"]) for s in summaries), reverse=True)
    rank_by_score = {score: ordered_scores.index(score) + 1 for score in ordered_scores}
    for summary in summaries:
        summary["순위"] = rank_by_score[float(summary["최종점수"])]
    summaries.sort(key=lambda x: (int(x["순위"]), str(x["피평가자"])))

    return summaries, details, errors


def style_sheet(sheet, widths: dict[int, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = TITLE_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")


def write_results(
    output_path: Path,
    summaries: list[dict[str, object]],
    details: list[dict[str, object]],
    errors: list[str],
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "순위"
    summary_sheet.append(["순위", "피평가자", "최종점수", "유효 평가자 수", "가중치 합"])
    for row in summaries:
        summary_sheet.append(
            [row["순위"], row["피평가자"], row["최종점수"], row["유효 평가자 수"], row["가중치 합"]]
        )
    style_sheet(summary_sheet, {1: 10, 2: 14, 3: 14, 4: 18, 5: 14})
    for cell in summary_sheet["C"][1:]:
        cell.number_format = "0.000"
    for cell in summary_sheet["E"][1:]:
        cell.number_format = "0.000"

    detail_sheet = workbook.create_sheet("평가자별 점수")
    detail_sheet.append(["피평가자", "평가자", "상피여부", "원점수", "가중치", "가중점수", "항목수"])
    for row in details:
        detail_sheet.append(
            [
                row["피평가자"],
                row["평가자"],
                row["상피여부"],
                row["원점수"],
                row["가중치"],
                row["가중점수"],
                row["항목수"],
            ]
        )
    style_sheet(detail_sheet, {1: 14, 2: 30, 3: 12, 4: 12, 5: 12, 6: 14, 7: 12})
    for cell in detail_sheet["E"][1:]:
        cell.number_format = "0.0"
    for cell in detail_sheet["F"][1:]:
        cell.number_format = "0.000"

    check_sheet = workbook.create_sheet("검증")
    check_sheet.append(["검증 결과", "내용"])
    if errors:
        for error in errors:
            check_sheet.append(["확인 필요", error])
    else:
        check_sheet.append(["정상", "평가자, 상피제, 평가 항목 수, 가중치 합 검증을 통과했습니다."])
    style_sheet(check_sheet, {1: 16, 2: 90})
    check_sheet["A2"].fill = SUBTITLE_FILL

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="개인정보 고래상 평가 결과 계산")
    parser.add_argument("input", type=Path, help="구글폼 응답 XLSX 파일")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="결과 XLSX 파일 경로(기본값: 입력파일명_평가결과.xlsx)",
    )
    args = parser.parse_args()

    input_path = args.input.resolve()
    output_path = args.output.resolve() if args.output else input_path.with_name(f"{input_path.stem}_평가결과.xlsx")

    headers, responses = read_latest_responses(input_path)
    summaries, details, errors = calculate_scores(headers, responses)
    write_results(output_path, summaries, details, errors)

    print(f"결과 파일: {output_path}")
    print(f"검증 결과: {'정상' if not errors else f'확인 필요 {len(errors)}건'}")
    for summary in summaries:
        print(f"{summary['순위']:>2}위  {summary['피평가자']}: {summary['최종점수']:.3f}점")

    if errors:
        for error in errors:
            print(f"- {error}")
        raise SystemExit(1)


if __name__ == "__main__":
    main()
